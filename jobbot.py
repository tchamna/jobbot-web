#!/usr/bin/env python3
"""
jobbot — tailored resume + cover letter generator and application tracker.

Usage
  python jobbot.py apply <job-url-or-.txt> [--company X --title Y] [--llm] [--no-pdf]
  python jobbot.py batch urls.txt              # one URL or .txt path per line
  python jobbot.py list                        # show tracker
  python jobbot.py status <id> <status>        # e.g. applied / interview / rejected / offer

What it does for each job
  1. Fetches the posting (or reads a pasted .txt), extracts title/company/text.
  2. Detects the job archetype (logistics, energy, healthcare, ai_llm, analytics).
  3. Scores every bullet in profile.yaml against the ad and keeps the strongest.
  4. Writes Resume + Cover Letter (.docx, plus .pdf when Word/LibreOffice is available).
  5. Writes a fit report: matched keywords, gaps to address, screening answers.
  6. Logs the job in tracker.csv.
It never submits applications for you; review the files, then apply.
"""
from __future__ import annotations
import argparse, csv, datetime as dt, json, os, re, shutil, subprocess, sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parent
PROFILE = ROOT / "profile.yaml"
OUT = ROOT / "output"
OVR = ROOT / "overrides"
TRACKER = ROOT / "tracker.csv"
ACCENT = RGBColor(0x1F, 0x4E, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)

# Vocabulary the bot looks for in ads (used for matching + gap report).
LEXICON = [
    "python", "sql", "r", "pyspark", "spark", "scala", "java", "c++", "matlab",
    "pandas", "numpy", "scikit-learn", "pytorch", "tensorflow", "xgboost",
    "machine learning", "deep learning", "statistical modeling", "statistics",
    "forecasting", "time series", "optimization", "linear programming",
    "mixed integer", "operations research", "simulation", "routing", "vrp",
    "scheduling", "inventory", "supply chain", "logistics", "warehouse", "wms",
    "tms", "erp", "sap", "anomaly detection", "a/b testing", "experimentation",
    "causal inference", "nlp", "llm", "rag", "generative ai", "computer vision",
    "azure", "aws", "gcp", "docker", "kubernetes", "airflow", "snowflake",
    "databricks", "dbt", "mlflow", "ci/cd", "git", "fastapi", "flask", "api",
    "power bi", "tableau", "dashboard", "etl", "data pipeline", "production",
    "deployment", "monitoring", "stakeholder", "communication", "gurobi",
    "or-tools", "pulp", "cplex", "master", "phd", "energy", "healthcare",
]

# ---------------------------------------------------------------- fetching
def fetch_job(src: str) -> dict:
    p = Path(src)
    if p.exists():
        raw = p.read_text(encoding="utf-8")
        meta = dict(re.findall(r"^(URL|TITLE|COMPANY|LOCATION|SALARY):\s*(.+)$", raw, re.M))
        return {"url": meta.get("URL", str(p)), "title": meta.get("TITLE", ""),
                "company": meta.get("COMPANY", ""), "location": meta.get("LOCATION", ""),
                "salary": meta.get("SALARY", ""), "text": raw}
    import requests
    from bs4 import BeautifulSoup
    r = requests.get(src, timeout=30, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                      "(KHTML, like Gecko) Chrome/126 Safari/537.36"})
    if r.status_code >= 400:
        raise RuntimeError(f"Could not read {src} (HTTP {r.status_code}). Many job boards block scripts.\n"
                 "Copy the posting text into a file, e.g. jobs/company_role.txt, with optional first lines\n"
                 "TITLE:, COMPANY:, LOCATION:, SALARY:, URL: and run:  python jobbot.py apply jobs/company_role.txt")
    soup = BeautifulSoup(r.text, "html.parser")
    job = {"url": src, "title": "", "company": "", "location": "", "salary": "", "text": ""}
    for tag in soup.find_all("script", type="application/ld+json"):   # schema.org JobPosting
        try:
            data = json.loads(tag.string or "")
        except Exception:
            continue
        for d in (data if isinstance(data, list) else [data]):
            if isinstance(d, dict) and d.get("@type") == "JobPosting":
                job["title"] = d.get("title", "")
                org = d.get("hiringOrganization") or {}
                job["company"] = org.get("name", "") if isinstance(org, dict) else str(org)
                job["text"] = BeautifulSoup(d.get("description", ""), "html.parser").get_text(" ")
    if not job["title"]:
        h1 = soup.find("h1"); h2 = soup.find("h2")
        job["title"] = h1.get_text(strip=True) if h1 else ""
        job["company"] = h2.get_text(strip=True) if h2 else ""
    if len(job["text"]) < 300:
        for s in soup(["script", "style", "nav", "footer", "header"]):
            s.decompose()
        job["text"] = soup.get_text("\n", strip=True)
    return job

# ---------------------------------------------------------------- scoring
def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.lower())

def has(term: str, text: str) -> bool:
    return re.search(r"(?<![a-z0-9])" + re.escape(term) + r"(?![a-z0-9])", text) is not None

def detect_archetype(profile: dict, jd: str) -> tuple[str, dict]:
    scores = {k: sum(jd.count(w) for w in a["detect"]) for k, a in profile["archetypes"].items()}
    best = max(scores, key=scores.get)
    return (best if scores[best] > 0 else "analytics"), scores

def score_item(text: str, tags: list, jd: str, jd_terms: set) -> float:
    s = sum(3 for t in tags if t.lower() in jd)                         # tag appears in ad
    s += sum(1 for t in jd_terms if has(t, norm(text)))                   # ad keyword in bullet
    s += 1 if re.search(r"\d+%|\d+\+|under \d", text) else 0             # quantified bonus
    return s

def select(profile: dict, job: dict):
    jd = norm(job["text"] + " " + job["title"])
    jd_terms = {t for t in LEXICON if has(t, jd)}
    arch, arch_scores = detect_archetype(profile, jd)
    exp = []
    for e in profile["experience"]:
        ranked = sorted(e["bullets"], key=lambda b: (not b.get("core"),
                        -score_item(b["text"], b.get("tags", []), jd, jd_terms)))
        keep = ranked[: e.get("max_bullets", 3)]
        keep.sort(key=lambda b: (not b.get("core"), -score_item(b["text"], b.get("tags", []), jd, jd_terms)))
        exp.append({**e, "bullets": keep})
    projects = sorted(profile.get("projects", []),
                      key=lambda p: -score_item(p["text"], p.get("tags", []), jd, jd_terms))[:3]
    projects = [p for p in projects if score_item(p["text"], p.get("tags", []), jd, jd_terms) > 0] or projects[:2]
    detect = profile["archetypes"][arch]["detect"]
    skills = sorted(profile["skills"], key=lambda g: -6 * any(t in detect for t in g["tags"])
                    - sum(3 for t in g["tags"] if t in jd)
                    - sum(1 for i in g["items"] if has(norm(i.split(" (")[0]), jd)))
    evidence = {k: v for k, v in profile.items() if k not in ("archetypes", "first_look_ideas")}
    profile_blob = norm(yaml.safe_dump(evidence))   # only real experience counts as a match
    matched = sorted(t for t in jd_terms if has(t, profile_blob))
    gaps = sorted(t for t in jd_terms if not has(t, profile_blob))
    return {"arch": arch, "arch_scores": arch_scores, "experience": exp, "projects": projects,
            "skills": skills, "matched": matched, "gaps": gaps, "jd": jd}

# ---------------------------------------------------------------- docx helpers
def base_doc() -> Document:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.65)
    sec.top_margin = sec.bottom_margin = Inches(0.55)
    st = d.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.05
    return d

def bottom_rule(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "1F4E5F"}.items():
        b.set(qn(k), v)
    bdr.append(b); pPr.append(bdr)

def heading(d, text):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = ACCENT
    r.font.all_caps = True
    bottom_rule(p)

def hyperlink(p, url, text, size=9.5):
    if not url.startswith("http"):
        url = "https://" + url
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1F4E5F"); rPr.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    r.append(rPr); t = OxmlElement("w:t"); t.text = text; r.append(t); h.append(r)
    p._p.append(h)

def bullet(d, text, bold_terms=()):
    p = d.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.left_indent = Inches(0.22); pf.first_line_indent = Inches(-0.16); pf.space_after = Pt(1.5)
    p.add_run(text)
    return p

def right_tab_line(d, left, right, left_bold=True, size=10.5, color=None):
    sec = d.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    p = d.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left); r.bold = left_bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    r2 = p.add_run("\t" + right); r2.font.size = Pt(size - 0.5); r2.font.color.rgb = GREY
    return p

def header_block(d, profile, headline):
    p = d.add_paragraph(); r = p.add_run(profile["name"]); r.bold = True
    r.font.size = Pt(21); r.font.color.rgb = ACCENT
    p = d.add_paragraph(); r = p.add_run(headline); r.font.size = Pt(11.5); r.bold = True
    r.font.color.rgb = GREY; p.paragraph_format.space_after = Pt(2)
    p = d.add_paragraph()
    r = p.add_run(f'{profile["location"]}  |  {profile["phone"]}  |  {profile["email"]}')
    r.font.size = Pt(9.5)
    p = d.add_paragraph()
    for i, l in enumerate(profile["links"]):
        if i: p.add_run("  |  ").font.size = Pt(9.5)
        hyperlink(p, l["url"], l["label"])
    return p

# ---------------------------------------------------------------- resume
def build_resume(profile, sel, path: Path):
    a = profile["archetypes"][sel["arch"]]
    d = base_doc()
    header_block(d, profile, sel.get("headline") or a["headline"])
    heading(d, "Summary")
    d.add_paragraph(sel.get("summary") or a["summary"])

    heading(d, "Core Skills")
    for g in sel["skills"][:5]:
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(1)
        r = p.add_run(g["group"] + ": "); r.bold = True
        p.add_run(", ".join(g["items"]))

    heading(d, "Experience")
    for e in sel["experience"]:
        p = right_tab_line(d, e["role"], e["dates"])
        p.paragraph_format.space_before = Pt(4)
        p = d.add_paragraph(); r = p.add_run(f'{e["org"]}  ·  {e["where"]}'); r.italic = True
        r.font.color.rgb = ACCENT; p.paragraph_format.space_after = Pt(1)
        for b in e["bullets"]:
            bullet(d, b["text"])

    heading(d, "Selected Projects")
    for pr in sel["projects"]:
        p = d.add_paragraph(style="List Bullet")
        pf = p.paragraph_format; pf.left_indent = Inches(0.22); pf.first_line_indent = Inches(-0.16)
        pf.space_after = Pt(1.5)
        r = p.add_run(f'{pr["name"]} ({pr["year"]}). '); r.bold = True
        p.add_run(pr["text"])
        if pr.get("url"):
            p.add_run(" ")
            hyperlink(p, pr["url"], pr["url"], 9.5)

    heading(d, "Education")
    for ed in profile["education"]:
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(1)
        r = p.add_run(ed["degree"]); r.bold = True
        p.add_run(f' — {ed["school"]}')
        if ed.get("note"):
            q = d.add_paragraph(); q.paragraph_format.left_indent = Inches(0.22)
            q.paragraph_format.space_after = Pt(1)
            rr = q.add_run(ed["note"]); rr.font.size = Pt(9.5); rr.italic = True

    heading(d, "Distinctions")
    for x in profile.get("distinctions", []):
        bullet(d, x)
    d.core_properties.author = profile["name"]
    d.core_properties.title = f'Resume — {profile["name"]}'
    d.save(path)

# ---------------------------------------------------------------- cover letter
def first_look(profile, jd):
    ideas = [f["idea"] for f in profile.get("first_look_ideas", []) if any(w in jd for w in f["when"])]
    if not ideas:
        return ""
    ideas = ideas[:3]
    listing = ideas[0] if len(ideas) == 1 else ", ".join(ideas[:-1]) + ", and " + ideas[-1]
    return ("If I joined, I would learn the operation before modeling it. The first places I would "
            f"look are {listing}.")

def proof_paragraphs(sel):
    jd = sel["jd"]
    paras = []
    optim = any(w in jd for w in ["optimiz", "routing", "logistic", "operations research", "schedul", "dispatch"])
    if optim:
        paras.append(
            "On optimization: I recently built a fleet-dispatch optimizer that models the problem as a "
            "pickup-and-delivery vehicle routing problem with time windows in Google OR-Tools, "
            "benchmarked it against greedy heuristics, and explained to non-specialists why the global "
            "solver makes better plans. At CUNY's transportation research center I modeled how "
            "social-distancing rules cut transit vehicle capacity, and I published and presented the "
            "results to public agencies.")
    paras.append(
        "On production: at EOS Energy I built the pipelines and real-time anomaly detection behind a "
        "grid-scale battery fleet, cutting processing time by 98% and response time by 75%, and "
        "deployed State-of-Charge models (under 3% RMSE) as FastAPI services in Docker. I ship to Azure "
        "with CI/CD, validate models against real outcomes, and prefer an interpretable model an "
        "operator trusts to a black box nobody uses.")
    return paras

def build_cover(profile, sel, job, ovr, path: Path):
    a = profile["archetypes"][sel["arch"]]
    company = ovr.get("company_name") or job["company"] or "your company"
    title = job["title"] or "this"
    fmt = lambda s: re.sub(r"\s+", " ", s.format(company=company, title=title)).strip()
    d = base_doc()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.9); s.top_margin = Inches(0.7)
    d.styles["Normal"].font.size = Pt(11)
    p = d.add_paragraph(); r = p.add_run(profile["name"]); r.bold = True; r.font.size = Pt(18)
    r.font.color.rgb = ACCENT
    p = d.add_paragraph(); r = p.add_run(f'{profile["location"]}  |  {profile["phone"]}  |  {profile["email"]}')
    r.font.size = Pt(9.5)
    p = d.add_paragraph()
    for i, l in enumerate(profile["links"][:1] + profile["links"][2:]):
        if i: p.add_run("  |  ").font.size = Pt(9.5)
        hyperlink(p, l["url"], l["label"])
    bottom_rule(p)
    gap = lambda: d.add_paragraph().paragraph_format.space_after or None
    blocks = [dt.date.today().strftime("%B %d, %Y").replace(" 0", " "),
              f"Re: {title}" + (f" ({job['location']})" if job.get("location") else ""),
              ovr.get("greeting") or f"Dear Hiring Team at {company},",
              fmt(ovr.get("hook") or a["hook"])]
    blocks += proof_paragraphs(sel)
    if ovr.get("why_company"):
        blocks.append(fmt(ovr["why_company"]))
    fl = ovr.get("first_look") or first_look(profile, sel["jd"])
    if fl:
        blocks.append(fmt(fl))
    blocks.append(fmt(ovr.get("closing") or
        f"I hold a PhD in Mechanical Engineering and a Master's in Physics, and I would welcome a "
        f"conversation about how I can contribute to {company}."))
    blocks += ["Thank you for your time and consideration.", "Sincerely,", profile["name"]]
    for i, b in enumerate(blocks):
        p = d.add_paragraph(b)
        p.paragraph_format.space_after = Pt(9 if i not in (len(blocks) - 2,) else 2)
        if i == 1: p.runs[0].bold = True
    d.core_properties.author = profile["name"]
    d.save(path)
    return blocks

# ---------------------------------------------------------------- optional LLM polish
def llm_polish(profile, job, sel, ovr):
    """If ANTHROPIC_API_KEY is set and --llm is passed, ask Claude for a company-specific
    'why_company' paragraph grounded ONLY in the ad and the profile."""
    try:
        import anthropic
    except ImportError:
        print("  (pip install anthropic to use --llm)"); return ovr
    if not os.environ.get("ANTHROPIC_API_KEY"):
        print("  (set ANTHROPIC_API_KEY to use --llm)"); return ovr
    client = anthropic.Anthropic()
    prompt = (
        "You write one paragraph (90-130 words) for a cover letter: why this candidate fits THIS "
        "company and role. Use only facts that appear in the PROFILE or the JOB AD. Do not invent "
        "metrics, employers, tools, or company facts. Plain, confident, first person, no cliches.\n\n"
        f"JOB AD:\n{job['text'][:8000]}\n\nPROFILE:\n{yaml.safe_dump(profile)[:12000]}")
    msg = client.messages.create(model=os.environ.get("JOBBOT_MODEL", "claude-opus-5"),
                                 max_tokens=400, messages=[{"role": "user", "content": prompt}])
    ovr = dict(ovr); ovr.setdefault("why_company", msg.content[0].text.strip())
    return ovr

# ---------------------------------------------------------------- pdf + report + tracker
def to_pdf(docx: Path) -> Path | None:
    try:                                    # Windows / macOS with Word installed
        from docx2pdf import convert
        convert(str(docx), str(docx.with_suffix(".pdf"))); return docx.with_suffix(".pdf")
    except Exception:
        pass
    exe = shutil.which("soffice") or shutil.which("libreoffice") or \
        next((p for p in [r"C:\Program Files\LibreOffice\program\soffice.exe"] if Path(p).exists()), None)
    if exe:
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        if docx.with_suffix(".pdf").exists():
            return docx.with_suffix(".pdf")
    return None

def write_report(job, sel, path: Path):
    jd = sel["jd"]
    lines = [f"# Fit report — {job['title']} @ {job['company']}", "",
             f"- URL: {job['url']}", f"- Location: {job.get('location','')}", f"- Salary: {job.get('salary','')}",
             f"- Archetype used: **{sel['arch']}**  (scores: {sel['arch_scores']})", "",
             f"## Keyword coverage: {len(sel['matched'])}/{len(sel['matched']) + len(sel['gaps'])}",
             "Matched: " + ", ".join(sel["matched"]), "",
             "Gaps (in the ad, not in your profile — address honestly or add to profile.yaml if true): "
             + (", ".join(sel["gaps"]) or "none"), "",
             "## Ready answers for screening questions",
             "- Highest degree: PhD, Mechanical Engineering (also M.S. Physics)",
             "- Years of applied data science / analytics: 8+ (2017–present), 10+ including research",
             "- Python: expert · SQL: advanced · Azure: yes (App Service, Blob) · Docker: yes",
             "- Work authorization / sponsorship: FILL IN",
             "- Salary expectation: " + (f"within posted range ({job['salary']})" if job.get("salary") else "FILL IN"),
             "", "## Before you submit",
             "- [ ] Read the resume top to bottom; delete anything you cannot talk about for 2 minutes",
             "- [ ] Check dates for NJ Sharing Network", "- [ ] Save as PDF, file name includes your name",
             "- [ ] Find the hiring manager on LinkedIn and send a 2-line note"]
    if "master" in jd and "2" in jd and "4 years" in jd:
        lines.insert(8, "> Note: the ad asks for 2–4 years. You exceed it; the letter frames depth as "
                        "hands-on applied work, not seniority, so you don't read as overqualified.")
    path.write_text("\n".join(lines), encoding="utf-8")

def log(job, folder: Path):
    new = not TRACKER.exists()
    rows = list(csv.DictReader(TRACKER.open(encoding="utf-8"))) if not new else []
    jid = str(len(rows) + 1)
    with TRACKER.open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["id", "date", "company", "title", "location", "salary", "url", "folder", "status"])
        if new: w.writeheader()
        w.writerow({"id": jid, "date": dt.date.today().isoformat(), "company": job["company"],
                    "title": job["title"], "location": job.get("location", ""), "salary": job.get("salary", ""),
                    "url": job["url"], "folder": str(folder.relative_to(ROOT)), "status": "prepared"})
    return jid

def slug(s: str) -> str:
    s = re.sub(r"[,.]?\s*(inc|llc|ltd|corp|co)\.?$", "", s.strip(), flags=re.I)
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")[:50] or "job"

# ---------------------------------------------------------------- commands
def cmd_apply(src, args):
    profile = yaml.safe_load(PROFILE.read_text(encoding="utf-8"))
    job = fetch_job(src)
    if args.company: job["company"] = args.company
    if args.title: job["title"] = args.title
    sel = select(profile, job)
    ovr_file = OVR / f"{slug(job['company'])}.yaml"
    ovr = yaml.safe_load(ovr_file.read_text(encoding="utf-8")) if ovr_file.exists() else {}
    ovr = ovr or {}
    if args.llm: ovr = llm_polish(profile, job, sel, ovr)
    sel["headline"] = ovr.get("headline"); sel["summary"] = ovr.get("summary")
    folder = OUT / f"{dt.date.today():%Y-%m-%d}_{slug(job['company'])}_{slug(job['title'])[:30]}"
    folder.mkdir(parents=True, exist_ok=True)
    last = profile["name"].split(",")[0].split()[-1]
    res = folder / f"Resume_{last}_{slug(job['company'])}.docx"
    cov = folder / f"CoverLetter_{last}_{slug(job['company'])}.docx"
    build_resume(profile, sel, res)
    build_cover(profile, sel, job, ovr, cov)
    write_report(job, sel, folder / "fit_report.md")
    if not args.no_pdf:
        for f in (res, cov):
            if not to_pdf(f): print(f"  (no PDF converter found for {f.name}; open in Word and Save as PDF)")
    jid = log(job, folder)
    print(f"[{jid}] {job['title']} @ {job['company']}  ->  {folder}")
    print(f"     archetype={sel['arch']}  matched={len(sel['matched'])}  gaps={sel['gaps']}")

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("apply"); a.add_argument("src")
    b = sub.add_parser("batch"); b.add_argument("file")
    for x in (a, b):
        x.add_argument("--company"); x.add_argument("--title")
        x.add_argument("--llm", action="store_true"); x.add_argument("--no-pdf", action="store_true")
    sub.add_parser("list")
    s = sub.add_parser("status"); s.add_argument("id"); s.add_argument("status")
    args = ap.parse_args()
    if args.cmd == "apply":
        cmd_apply(args.src, args)
    elif args.cmd == "batch":
        for line in Path(args.file).read_text(encoding="utf-8").splitlines():
            if line.strip() and not line.startswith("#"):
                try: cmd_apply(line.strip(), args)
                except Exception as e: print(f"  FAILED {line.strip()}: {e}")
    elif args.cmd == "list":
        if not TRACKER.exists(): print("No applications yet."); return
        for r in csv.DictReader(TRACKER.open(encoding="utf-8")):
            print(f'{r["id"]:>3}  {r["date"]}  {r["status"]:<10} {r["company"][:28]:<28} {r["title"][:45]}')
    elif args.cmd == "status":
        rows = list(csv.DictReader(TRACKER.open(encoding="utf-8")))
        for r in rows:
            if r["id"] == args.id: r["status"] = args.status
        with TRACKER.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=rows[0].keys()); w.writeheader(); w.writerows(rows)
        print("updated")

if __name__ == "__main__":
    main()
